"""Vietnamese legislation pipeline (R4). Synthetic statute: "FIXTURE — không phải luật"."""

from __future__ import annotations

import io
import unicodedata
from datetime import date
from pathlib import Path

import pytest
from travo_rag.sources.fetch import SnapshotStore
from travo_rag.sources.manifest import Instrument, Manifest, load_manifests
from travo_rag.sources.parsers import ParseError, parse_snapshot
from travo_rag.sources.pipeline import run, write_outputs

ROOT = Path(__file__).resolve().parents[1]

LINES = [
    "QUỐC HỘI",
    "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM",
    "Độc lập - Tự do - Hạnh phúc",
    "Luật số: 99/2099/QH99",
    "LUẬT FIXTURE HỢP ĐỒNG (FIXTURE — không phải luật)",
    "Căn cứ Hiến pháp FIXTURE;",
    "Chương I",
    "NHỮNG QUY ĐỊNH CHUNG",
    "Điều 1. Phạm vi điều chỉnh",
    "Luật FIXTURE này quy định về hợp đồng FIXTURE giữa các bên.",
    "Điều 2. Phạt vi phạm",
    "1. Mức phạt vi phạm FIXTURE do các bên thỏa thuận trong hợp đồng, nhưng tổng mức phạt",
    "không vượt quá 8% giá trị phần nghĩa vụ bị vi phạm, trừ trường hợp quy định tại",
    "Điều 3 của Luật FIXTURE này.",
    "2. Bên vi phạm phải chịu phạt trong các trường hợp sau đây:",
    "a) Không thực hiện nghĩa vụ FIXTURE;",
    "b) Thực hiện không đúng nghĩa vụ FIXTURE.",
    "[1] Khoản này được sửa đổi bởi Luật FIXTURE số 98/2098/QH98.",
    "Chương II",
    "ĐIỀU KHOẢN THI HÀNH",
    "Điều 3. Trường hợp đặc biệt",
    "Các bên FIXTURE có thể thỏa thuận khác trong trường hợp đặc biệt.",
    "Điều 4. Hiệu lực thi hành",
    "Luật FIXTURE này có hiệu lực thi hành từ ngày 01 tháng 07 năm 2020.",
    "Luật này đã được Quốc hội FIXTURE thông qua ngày 01 tháng 01 năm 2020.",
    "CHỦ TỊCH QUỐC HỘI FIXTURE",
]


def inst(**kw) -> Instrument:
    base = {
        "id": "VN/FIXTURE99",
        "title": "Luật FIXTURE Hợp đồng 2099",
        "instrument_type": "law",
        "number": "99/2099/QH99",
        "url": None,
        "format": "vbpl_html",
        "language": "vi",
        "expect_title": "Luật FIXTURE Hợp đồng",
    }
    return Instrument.model_validate({**base, **kw})


def html(lines: list[str]) -> bytes:
    body = "".join(f"<p>{ln}</p>" for ln in lines)
    return (
        "<html><head><title>FIXTURE vbpl</title></head><body><div class='menu'>Trang chủ</div>"
        f"<div id='toanvancontent'>{body}</div><div class='footer'>© FIXTURE</div></body></html>"
    ).encode()


def docx(lines: list[str]) -> bytes:
    from docx import Document

    d = Document()
    t = d.add_table(rows=1, cols=2)  # header block: issuing body | national motto
    t.rows[0].cells[0].text = lines[0]
    t.rows[0].cells[1].text = f"{lines[1]}\n{lines[2]}"
    for ln in lines[3:]:
        d.add_paragraph(ln)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _parse(tmp_path, fmt: str, data: bytes):
    i = inst(format=fmt)
    snap = SnapshotStore(tmp_path).save(i, None, data, "x", origin="supplied", filename=f"a.{fmt}")
    return parse_snapshot(i, snap, "FIXTURE")


@pytest.mark.parametrize("fmt", ["vbpl_html", "docx"])
def test_articles_parsed_with_khoan_diem_parts_notes(tmp_path, fmt):
    data = html(LINES) if fmt == "vbpl_html" else docx(LINES)
    r = _parse(tmp_path, fmt, data)
    assert [u.unit_path for u in r.units] == ["Điều 1", "Điều 2", "Điều 3", "Điều 4"]
    d2 = r.units[1]
    assert d2.heading == "Phạt vi phạm"
    lines = d2.text.splitlines()
    assert lines[0].startswith("1. Mức phạt") and lines[0].endswith("Điều 3 của Luật FIXTURE này.")
    assert lines[1:] == [
        "2. Bên vi phạm phải chịu phạt trong các trường hợp sau đây:",
        "a) Không thực hiện nghĩa vụ FIXTURE;",
        "b) Thực hiện không đúng nghĩa vụ FIXTURE.",
    ]
    assert r.parts["Điều 1"] == "Chương I NHỮNG QUY ĐỊNH CHUNG"
    assert r.parts["Điều 3"] == "Chương II ĐIỀU KHOẢN THI HÀNH"
    assert r.notes["Điều 2"][0].startswith("[1] Khoản này")
    assert "đã được Quốc hội" not in r.units[-1].text  # closing formula is not an article
    assert all(u.effective_from == date(2020, 7, 1) for u in r.units)
    assert r.stats["document number"] == "99/2099/QH99"
    assert not [w for w in r.warnings if "accounting" in w], r.warnings
    assert all(u.source.language == "vi" for u in r.units)


def test_decomposed_unicode_is_normalised(tmp_path):
    nfd = [unicodedata.normalize("NFD", ln) for ln in LINES]
    r = _parse(tmp_path, "vbpl_html", html(nfd))
    assert r.units[1].heading == unicodedata.normalize("NFC", "Phạt vi phạm")


def test_legacy_font_encoding_is_refused(tmp_path):
    tcvn3 = ["Qu¸c héi", "Lu©t sè 99/2099/QH99", "Lu©t FIXTURE Hîp ®ång"] + [
        "§iÒu 1. Ph¹m vi ®iÒu chØnh vµ ®èi t­îng ¸p dông cña luËt nµy"
    ] * 20
    with pytest.raises(ParseError, match="legacy Vietnamese font encoding"):
        _parse(tmp_path, "vbpl_html", html(tcvn3))


def test_quoted_article_reference_does_not_start_an_article(tmp_path):
    lines = [*LINES[:11], "Điều 4. Tham chiếu", *LINES[11:]]  # out of order: stays text
    r = _parse(tmp_path, "vbpl_html", html(lines))
    assert [u.unit_path for u in r.units] == ["Điều 1", "Điều 2", "Điều 3", "Điều 4"]
    assert "Điều 4. Tham chiếu" in r.units[1].text


def test_vn_manifest_loads():
    m = load_manifests(ROOT / "config/legal_sources")["VN"]
    assert len(m.instruments) == 15
    assert all(i.language == "vi" and i.format == "vbpl_html" for i in m.instruments)


def test_ingest_search_without_tone_marks_and_vietnamese_pinpoints(tmp_path, client, make_tenant):
    from travo_api.db import get_admin_engine, tenant_session
    from travo_rag import citations as cite
    from travo_rag import retrieval
    from travo_rag.legal_index import upsert_units

    r = _parse(tmp_path, "vbpl_html", html(LINES))
    with get_admin_engine().begin() as c:
        upsert_units(c, r.units)
    t = make_tenant()
    with tenant_session(t.tenant_id) as s:
        for query in ("phạt vi phạm vượt quá giá trị", "phat vi pham vuot qua gia tri"):
            hits = retrieval.search(s, query, jurisdictions=["VN"], as_of=date(2026, 9, 30), k=3)
            assert hits and hits[0].unit_path == "Điều 2", (query, hits)
        hit = hits[0]
        assert hit.pinpoint == "Điều 2 Luật FIXTURE Hợp đồng 2099"
        # A quote in decomposed Unicode still matches the NFC source.
        quote = unicodedata.normalize("NFD", "không vượt quá 8% giá trị phần nghĩa vụ bị vi phạm")
        note = f'Theo luật, mức phạt "{quote}" [[src:{hit.id}]]. Điều này áp dụng.'
        results = cite.validate(
            note,
            lookup=lambda uid: retrieval.get_unit(s, uid),
            as_of=date(2026, 9, 30),
            require_cites=False,
        )
        assert [x.status for x in results] == ["supported"]


def test_pipeline_offline_report_for_vn(tmp_path):
    store = SnapshotStore(tmp_path / "snaps")
    i = inst()
    store.save(i, None, html(LINES), "text/html", origin="supplied", filename="luat.html")
    m = Manifest(jurisdiction="VN", issuing_body="FIXTURE", instruments=[i])
    outcomes = run(m, store, None)
    assert outcomes[0].status == "parsed" and len(outcomes[0].units) == 4
    _, report, n = write_outputs(outcomes, "VN", tmp_path / "out")
    text = report.read_text()
    assert n == 4 and "Điều 2" in text and "Articles parsed: 4" in text
