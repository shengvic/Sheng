"""Export gate + DOCX outputs (docs/05 §6-7): redline with Word tracked changes, review memo."""

from __future__ import annotations

import hashlib
import io
import itertools
import uuid
from datetime import UTC, datetime
from typing import Any

from docx import Document as DocxDocument
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from sqlalchemy import select
from travo_agents.taxonomy import CLAUSE_NAMES_VI

from travo_api.audit import audit
from travo_api.auth import Actor
from travo_api.ingest import get_store
from travo_api.keys import tenant_dek
from travo_api.models import Citation, Clause, Document, Export, Finding, ReviewRun
from travo_api.telemetry import emit

AUTHOR = "Travo (draft)"
PASSING_CITATION = {"supported"}


def blocking_items(actor: Actor, run: ReviewRun) -> list[dict[str, Any]]:
    """Why a review cannot be exported yet. Empty list = gate open."""
    blocks: list[dict[str, Any]] = []
    findings = list(actor.session.scalars(select(Finding).where(Finding.run_id == run.id)))
    for f in findings:
        if f.severity != "info" and f.disposition in (None, "deferred"):
            blocks.append(
                {
                    "type": "finding_undispositioned",
                    "finding_id": str(f.id),
                    "severity": f.severity,
                    "summary": f.summary[:200],
                }
            )
    live = {f.id for f in findings if f.disposition != "rejected"}
    if live:
        for c in actor.session.scalars(select(Citation).where(Citation.finding_id.in_(live))):
            if c.status not in PASSING_CITATION and not c.override_reason:
                blocks.append(
                    {
                        "type": "citation_unsupported",
                        "citation_id": str(c.id),
                        "finding_id": str(c.finding_id),
                        "status": c.status,
                    }
                )
    return blocks


class _Rev:
    """Monotonic ids for w:ins / w:del revision marks."""

    def __init__(self) -> None:
        self._ids = itertools.count(1)
        self.date = datetime.now(UTC).strftime("%Y-%m-%dT%H:%M:%SZ")

    def mark(self, tag: str) -> Any:
        el = OxmlElement(tag)
        el.set(qn("w:id"), str(next(self._ids)))
        el.set(qn("w:author"), AUTHOR)
        el.set(qn("w:date"), self.date)
        return el


def _run(text: str, deleted: bool = False) -> Any:
    r = OxmlElement("w:r")
    t = OxmlElement("w:delText" if deleted else "w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    r.append(t)
    return r


def _tracked_replace(p: Paragraph, old: str, new: str, rev: _Rev) -> None:
    if old:
        d = rev.mark("w:del")
        d.append(_run(old, deleted=True))
        p._p.append(d)
    ins = rev.mark("w:ins")
    ins.append(_run(new))
    p._p.append(ins)


def _bilingual(clauses: list[Clause]) -> bool:
    return any(c.text_alt for c in clauses)


def build_redline_docx(clauses: list[Clause], findings: list[Finding], title: str) -> bytes:
    """Rebuilds the contract from extracted clauses; accepted/edited redlines appear as Word
    tracked changes so the lawyer can accept or reject them in Word. Bilingual contracts are
    rebuilt as a two-column table (primary | other language), with the change in each column."""
    doc = DocxDocument()
    doc.add_heading(title, level=1)
    rev = _Rev()
    replacement: dict[uuid.UUID, tuple[str, str | None]] = {}
    appended: list[tuple[str, str, str | None]] = []
    for f in findings:
        if f.disposition not in ("accepted", "edited") or f.kind != "playbook":
            continue
        new = f.edited_text if f.disposition == "edited" else f.suggested_redline
        if not new:
            continue
        # An edited redline replaces the lawyer's language only; the other version is theirs
        # to align (the edit exists in one language).
        alt = f.suggested_redline_alt if f.disposition == "accepted" else None
        if f.clause_id:
            replacement[f.clause_id] = (new, alt)
        else:
            appended.append((f.clause_key.replace("_", " ").title(), new, alt))

    if not _bilingual(clauses):
        for c in clauses:
            heading = " ".join(x for x in (c.number, c.heading) if x)
            if heading:
                doc.add_paragraph(heading).runs[0].bold = True
            p = doc.add_paragraph()
            if c.id in replacement:
                _tracked_replace(p, c.text, replacement[c.id][0], rev)
            else:
                p.add_run(c.text)
        for heading, new, _ in appended:
            hp = doc.add_paragraph()
            _tracked_replace(hp, "", heading, rev)
            _tracked_replace(doc.add_paragraph(), "", new, rev)
    else:
        table = doc.add_table(rows=0, cols=2)
        table.style = "Table Grid"
        for c in clauses:
            left, right = table.add_row().cells
            heading = " ".join(x for x in (c.number, c.heading) if x)
            for cell, head in ((left, heading), (right, c.heading_alt)):
                if head:
                    cell.paragraphs[0].add_run(head).bold = True
                    cell.add_paragraph()
            new, alt = replacement.get(c.id, (None, None))
            p_left = left.paragraphs[-1]
            p_right = right.paragraphs[-1]
            if new is not None:
                _tracked_replace(p_left, c.text, new, rev)
            else:
                p_left.add_run(c.text)
            if alt is not None:
                _tracked_replace(p_right, c.text_alt, alt, rev)
            else:
                p_right.add_run(c.text_alt)
        for heading, new, alt in appended:
            left, right = table.add_row().cells
            _tracked_replace(left.paragraphs[0], "", heading, rev)
            _tracked_replace(left.add_paragraph(), "", new, rev)
            if alt:
                _tracked_replace(right.paragraphs[0], "", alt, rev)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


MEMO_LABELS: dict[str, dict[str, str]] = {
    "en": {
        "title": "Contract review memo — {name}",
        "meta": "Playbook: {pb} v{v} ({src}). Contract type: {ct}. Governing law: {law}.",
        "draft": "Draft prepared with Travo. Reviewed findings only; verify before sending.",
        "summary": "Executive summary",
        "summary_en": "Executive summary (English)",
        "points": "Negotiation points",
        "issues": "Issues list",
        "cols": "Clause|Severity|Finding|Disposition|Sources",
        "overridden": "overridden",
    },
    "vi": {
        "title": "Bản ghi nhớ rà soát hợp đồng — {name}",
        "meta": "Bộ quy tắc: {pb} v{v} ({src}). Loại hợp đồng: {ct}. Luật áp dụng: {law}.",
        "draft": (
            "Bản nháp do Travo hỗ trợ soạn. Chỉ gồm các phát hiện đã xem xét; "
            "kiểm tra trước khi gửi."
        ),
        "summary": "Tóm tắt",
        "summary_en": "Tóm tắt (tiếng Anh)",
        "points": "Điểm cần đàm phán",
        "issues": "Danh sách vấn đề",
        "cols": "Điều khoản|Mức độ|Phát hiện|Xử lý|Nguồn",
        "overridden": "đã bỏ qua kiểm tra",
    },
}
_VI_TERMS = {
    "high": "cao", "medium": "trung bình", "low": "thấp",
    "accepted": "đã chấp nhận", "edited": "đã sửa", "deferred": "để sau",
}  # fmt: skip


def build_memo_docx(
    run: ReviewRun,
    document: Document,
    findings: list[Finding],
    citations: dict[uuid.UUID, list[Citation]],
) -> bytes:
    out = (run.options or {}).get("output_language", "en")
    lang = "vi" if out in ("vi", "both") else "en"
    lab = MEMO_LABELS[lang]

    def term(x: str | None) -> str:
        return (_VI_TERMS.get(x or "", x or "") if lang == "vi" else (x or "")) or ""

    def clause(key: str) -> str:
        return CLAUSE_NAMES_VI.get(key, key) if lang == "vi" else key.replace("_", " ")

    doc = DocxDocument()
    doc.add_heading(lab["title"].format(name=document.filename), level=1)
    doc.add_paragraph(
        lab["meta"].format(
            pb=run.playbook_key,
            v=run.playbook_version,
            src=run.playbook_source,
            ct=document.contract_type,
            law=document.governing_law,
        )
    )
    doc.add_paragraph(lab["draft"])
    doc.add_heading(lab["summary"], level=2)
    doc.add_paragraph(str(run.summary.get("executive_summary", "")))
    if run.summary.get("executive_summary_en"):
        doc.add_heading(lab["summary_en"], level=2)
        doc.add_paragraph(str(run.summary["executive_summary_en"]))
    points = run.summary.get("negotiation_points") or []
    if points:
        doc.add_heading(lab["points"], level=2)
        for pt in points:
            doc.add_paragraph(str(pt), style="List Bullet")
    doc.add_heading(lab["issues"], level=2)
    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    for cell, h in zip(table.rows[0].cells, lab["cols"].split("|"), strict=True):
        cell.text = h
    order = {"high": 0, "medium": 1, "low": 2, "info": 3}
    for f in sorted(findings, key=lambda x: (order.get(x.severity, 9), x.clause_key)):
        if f.disposition == "rejected" or f.severity == "info":
            continue
        row = table.add_row().cells
        row[0].text = clause(f.clause_key)
        row[1].text = term(f.severity)
        row[2].text = f.edited_text if f.disposition == "edited" and f.kind == "law" else f.summary
        row[3].text = term(f.disposition)
        srcs = citations.get(f.id, [])
        row[4].text = "; ".join(
            f"{c.pinpoint or c.source_unit_id} ({c.status}"
            + (f", {lab['overridden']}" if c.override_reason else "")
            + ")"
            for c in srcs
        )
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def create_export(actor: Actor, run: ReviewRun, fmt: str) -> Export:
    s = actor.session
    document = s.get(Document, run.document_id)
    assert document is not None
    findings = list(s.scalars(select(Finding).where(Finding.run_id == run.id)))
    if fmt == "redline_docx":
        clauses = list(
            s.scalars(
                select(Clause).where(Clause.document_id == run.document_id).order_by(Clause.idx)
            )
        )
        data = build_redline_docx(clauses, findings, f"{document.filename} — Travo redline")
        filename = f"{document.filename.rsplit('.', 1)[0]}-redline.docx"
    else:
        cits: dict[uuid.UUID, list[Citation]] = {}
        if findings:
            for c in s.scalars(
                select(Citation).where(Citation.finding_id.in_([f.id for f in findings]))
            ):
                cits.setdefault(c.finding_id, []).append(c)
        data = build_memo_docx(run, document, findings, cits)
        filename = f"{document.filename.rsplit('.', 1)[0]}-review-memo.docx"
    ref = get_store().put(actor.tenant_id, tenant_dek(s, actor.tenant_id), data)
    export = Export(
        tenant_id=run.tenant_id,
        matter_id=run.matter_id,
        run_id=run.id,
        format=fmt,
        filename=filename[:300],
        storage_ref=ref,
        sha256=hashlib.sha256(data).hexdigest(),
        created_by=actor.user_id,
    )
    s.add(export)
    s.flush()
    audit(
        s,
        tenant_id=actor.tenant_id,
        actor_id=actor.user_id,
        action="export.created",
        resource_type="export",
        resource_id=export.id,
        matter_id=run.matter_id,
        details={"format": fmt, "run_id": str(run.id), "sha256": export.sha256},
    )
    emit(
        s,
        tenant_id=actor.tenant_id,
        event_type="export.created",
        matter_id=run.matter_id,
        actor_id=actor.user_id,
        subject_id=export.id,
        payload={"format": fmt},
    )
    return export


def read_export(actor: Actor, export: Export) -> bytes:
    return get_store().get(
        actor.tenant_id, tenant_dek(actor.session, actor.tenant_id), export.storage_ref
    )
