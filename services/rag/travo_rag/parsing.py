"""Parse uploaded contracts into ordered text blocks."""

from __future__ import annotations

import io
import re
from dataclasses import dataclass

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PDF_MIME = "application/pdf"
TEXT_MIME = "text/plain"
SUPPORTED_MIME = {DOCX_MIME, PDF_MIME, TEXT_MIME}


class UnsupportedDocument(ValueError):
    pass


@dataclass(frozen=True)
class Block:
    text: str
    is_heading: bool = False
    # Table rows keep their cells: bilingual contracts are often a VI | EN two-column table.
    cells: tuple[str, ...] = ()
    # The other-language text of this block, set by `travo_rag.bilingual.separate`.
    alt: str = ""


def sniff_mime(filename: str, declared: str | None, head: bytes) -> str:
    if head.startswith(b"%PDF"):
        return PDF_MIME
    if head.startswith(b"PK") and filename.lower().endswith(".docx"):
        return DOCX_MIME
    if declared in SUPPORTED_MIME:
        return declared
    if filename.lower().endswith(".txt"):
        return TEXT_MIME
    raise UnsupportedDocument(f"unsupported file type: {filename}")


def parse(data: bytes, mime: str) -> list[Block]:
    if mime == DOCX_MIME:
        return _parse_docx(data)
    if mime == PDF_MIME:
        return _parse_pdf(data)
    if mime == TEXT_MIME:
        return _lines(data.decode("utf-8", errors="replace"))
    raise UnsupportedDocument(mime)


def _parse_docx(data: bytes) -> list[Block]:
    from docx import Document  # lazy: heavy import
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = Document(io.BytesIO(data))
    blocks: list[Block] = []
    # Body order matters: a clause laid out as a table sits between paragraphs.
    for child in doc.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            p = Paragraph(child, doc)
            text = p.text.strip()
            if not text:
                continue
            style = (p.style.name or "") if p.style is not None else ""
            # Document titles belong to the preamble; only section headings start clauses.
            blocks.append(Block(text=text, is_heading=style.lower().startswith("heading")))
        elif tag == "tbl":
            for row in Table(child, doc).rows:
                seen: set[int] = set()
                cells: list[str] = []
                for cell in row.cells:
                    if id(cell._tc) in seen:  # merged cells repeat in python-docx
                        continue
                    seen.add(id(cell._tc))
                    if cell.text.strip():
                        cells.append(cell.text.strip())
                if cells:
                    blocks.append(Block(text=" | ".join(cells), cells=tuple(cells)))
    return blocks


def _parse_pdf(data: bytes) -> list[Block]:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    text = "\n".join((page.extract_text() or "") for page in reader.pages)
    if not text.strip():
        # Scanned PDF: OCR lands in P1 (docs/10).
        raise UnsupportedDocument("PDF has no text layer (OCR not yet supported)")
    return _lines(text)


def _lines(text: str) -> list[Block]:
    return [Block(text=ln.strip()) for ln in text.splitlines() if ln.strip()]


VI_CHARS = re.compile(r"[ăâđêôơưạảấầẩẫậắằẳẵặẹẻẽếềểễệỉịọỏốồổỗộớờởỡợụủứừửữựỳỵỷỹ]", re.I)
_ID_WORDS = {"yang", "dan", "dengan", "perjanjian", "pihak", "untuk", "dalam", "tersebut", "pasal"}
_MS_WORDS = {"hendaklah", "perjanjian", "pihak", "yang", "dan", "adalah", "fasal", "boleh"}
_EN_WORDS = {"the", "and", "of", "agreement", "shall", "party", "to", "in"}


def detect_languages(text: str, min_share: float = 0.15) -> list[str]:
    """Cheap heuristic language ID (en/vi/id/ms). Replaced by a T0 model in P1."""
    words = re.findall(r"[^\W\d_]+", text.lower())
    if not words:
        return []
    vi_words = sum(1 for w in words if VI_CHARS.search(w))
    # Latin-script languages are scored against non-Vietnamese words so bilingual
    # documents surface both languages.
    latin = max(len(words) - vi_words, 1)
    scores: dict[str, float] = {
        "vi": vi_words / len(words),
        "en": sum(1 for w in words if w in _EN_WORDS) / latin * 3,
    }
    id_share = sum(1 for w in words if w in _ID_WORDS) / latin * 3
    ms_share = sum(1 for w in words if w in _MS_WORDS) / latin * 3
    if id_share or ms_share:
        # Malay vs Indonesian share vocabulary; distinctive markers decide.
        if {"hendaklah", "fasal"} & set(words):
            scores["ms"] = ms_share
        else:
            scores["id"] = id_share
    langs = [k for k, v in sorted(scores.items(), key=lambda kv: -kv[1]) if v >= min_share]
    return langs or [max(scores, key=lambda k: scores[k])]
