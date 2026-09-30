"""Synthetic Vietnamese / bilingual contracts for tests and evals (ADR-023).

Everything here is invented: parties are FIXTURE companies and the terms are not advice. The
builders produce the four bilingual layouts Travo must read: a VI | EN two-column table, inline
pairs, alternating paragraphs and two halves; plus Vietnamese-only.
Clause tuples are (vi_heading, vi_body, en_heading, en_body, taxonomy_key).
"""

from __future__ import annotations

import io
from collections.abc import Sequence

from docx import Document

Clause = tuple[str, str, str, str, str]

NDA_CLAUSES: list[Clause] = [
    (
        "Định nghĩa",
        "Thông tin mật là mọi thông tin do Bên Tiết Lộ cung cấp cho Bên Nhận.",
        "Definitions",
        "Confidential Information means any information disclosed by the Discloser.",
        "definitions",
    ),
    (
        "Nghĩa vụ bảo mật",
        "Bên Nhận phải giữ bí mật và không tiết lộ Thông tin mật cho bên thứ ba.",
        "Confidentiality",
        "The Recipient shall keep the Confidential Information secret and shall not "
        "disclose it to any third party.",
        "confidentiality_obligations",
    ),
    (
        "Thời hạn",
        "Thỏa thuận này có hiệu lực trong thời hạn hai (02) năm kể từ ngày 01 tháng 02 năm 2026.",
        "Term",
        "This Agreement remains in force for two (2) years from 1 February 2026.",
        "term_and_termination",
    ),
    (
        "Phạt vi phạm",
        "Bên vi phạm phải trả tiền phạt 100.000.000 đồng (Bằng chữ: Một trăm triệu đồng).",
        "Penalty",
        "The breaching party shall pay a penalty of VND 100,000,000 (one hundred million dong).",
        "penalty",
    ),
    (
        "Luật áp dụng",
        "Thỏa thuận này được điều chỉnh bởi pháp luật Việt Nam.",
        "Governing Law",
        "This Agreement is governed by the laws of Vietnam.",
        "governing_law",
    ),
    (
        "Giải quyết tranh chấp",
        "Tranh chấp được giải quyết tại Trung tâm Trọng tài Quốc tế Việt Nam (VIAC).",
        "Dispute Resolution",
        "Disputes shall be settled by arbitration at the Vietnam International Arbitration "
        "Centre (VIAC).",
        "dispute_resolution",
    ),
    (
        "Ngôn ngữ",
        "Thỏa thuận được lập bằng tiếng Việt và tiếng Anh; bản tiếng Việt được ưu tiên áp dụng.",
        "Language",
        "This Agreement is made in Vietnamese and English; the Vietnamese version prevails.",
        "language",
    ),
]

TITLE = ("THỎA THUẬN BẢO MẬT", "NON-DISCLOSURE AGREEMENT")
PARTIES = (
    "Các bên: Công ty Cổ phần FIXTURE Sông Hồng và Công ty TNHH FIXTURE Sài Gòn.",
    "Parties: FIXTURE Song Hong JSC and FIXTURE Saigon Co., Ltd.",
)


def _save(doc) -> bytes:  # type: ignore[no-untyped-def]
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def table_docx(clauses: Sequence[Clause] = NDA_CLAUSES, trailer: str | None = None) -> bytes:
    """Title and parties as paragraphs, then one table row per clause (VI | EN)."""
    doc = Document()
    doc.add_heading(f"{TITLE[0]} / {TITLE[1]}", level=0)
    doc.add_paragraph(PARTIES[0])
    doc.add_paragraph(PARTIES[1])
    table = doc.add_table(rows=0, cols=2)
    for n, (vh, vb, eh, eb, _) in enumerate(clauses, 1):
        row = table.add_row().cells
        row[0].text = f"Điều {n}. {vh}\n{vb}"
        row[1].text = f"Article {n}. {eh}\n{eb}"
    if trailer:
        doc.add_paragraph(trailer)
    return _save(doc)


def inline_docx(clauses: Sequence[Clause] = NDA_CLAUSES) -> bytes:
    """ "Điều 1. Định nghĩa / Definitions" headings; VI sentence then EN sentence."""
    doc = Document()
    doc.add_heading(f"{TITLE[0]} / {TITLE[1]}", level=0)
    doc.add_paragraph(f"{PARTIES[0]} {PARTIES[1]}")
    for n, (vh, vb, eh, eb, _) in enumerate(clauses, 1):
        doc.add_paragraph(f"Điều {n}. {vh} / {eh}")
        doc.add_paragraph(f"{vb} {eb}")
    return _save(doc)


def paragraphs_docx(clauses: Sequence[Clause] = NDA_CLAUSES) -> bytes:
    """Each VI paragraph followed by its EN paragraph."""
    doc = Document()
    doc.add_heading(TITLE[0], level=0)
    doc.add_paragraph(TITLE[1])
    doc.add_paragraph(PARTIES[0])
    doc.add_paragraph(PARTIES[1])
    for n, (vh, vb, eh, eb, _) in enumerate(clauses, 1):
        doc.add_paragraph(f"Điều {n}. {vh}")
        doc.add_paragraph(f"Article {n}. {eh}")
        doc.add_paragraph(vb)
        doc.add_paragraph(eb)
    return _save(doc)


def halves_docx(clauses: Sequence[Clause] = NDA_CLAUSES) -> bytes:
    """The whole Vietnamese text, then the whole English text."""
    doc = Document()
    doc.add_heading(TITLE[0], level=0)
    doc.add_paragraph(PARTIES[0])
    for n, (vh, vb, _, _, _) in enumerate(clauses, 1):
        doc.add_paragraph(f"Điều {n}. {vh}")
        doc.add_paragraph(vb)
    doc.add_paragraph(TITLE[1])
    doc.add_paragraph(PARTIES[1])
    for n, (_, _, eh, eb, _) in enumerate(clauses, 1):
        doc.add_paragraph(f"Article {n}. {eh}")
        doc.add_paragraph(eb)
    return _save(doc)


def vietnamese_docx(clauses: Sequence[Clause] = NDA_CLAUSES) -> bytes:
    """Vietnamese only, with numbered sub-clauses (khoản) inside articles."""
    doc = Document()
    doc.add_heading(TITLE[0], level=0)
    doc.add_paragraph(PARTIES[0])
    for n, (vh, vb, _, _, _) in enumerate(clauses, 1):
        doc.add_paragraph(f"Điều {n}. {vh}")
        doc.add_paragraph(f"1. {vb}")
        doc.add_paragraph("2. Các bên cam kết thực hiện khoản này một cách thiện chí.")
    return _save(doc)


# Seeded discrepancies (for tests and the eval set): what a reviewer must catch.
SEEDED: dict[str, str] = {
    "term_and_termination": "duration",  # 2 years (VI) vs 3 years (EN)
    "penalty": "figure_words",  # 200.000.000 in figures, "Một trăm triệu" in words
    "confidentiality_obligations": "negation",  # VI forbids disclosure, EN does not say so
    "language": "prevailing_language",  # each version says it prevails
}


def seeded_clauses() -> list[Clause]:
    out = []
    for vh, vb, eh, eb, key in NDA_CLAUSES:
        if key == "term_and_termination":
            eb = "This Agreement remains in force for three (3) years from 1 February 2026."
        elif key == "penalty":
            vb = "Bên vi phạm phải trả tiền phạt 200.000.000 đồng (Bằng chữ: Một trăm triệu đồng)."
            eb = "The breaching party shall pay a penalty of VND 200,000,000."
        elif key == "confidentiality_obligations":
            eb = "The Recipient shall keep the Confidential Information secret."
        elif key == "language":
            eb = "This Agreement is made in Vietnamese and English; the English version prevails."
        out.append((vh, vb, eh, eb, key))
    return out
